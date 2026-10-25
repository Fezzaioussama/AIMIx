"""Excel workbooks and PowerPoint decks: what text they give, the limits on them, and
the package guards they share with Word, with what a hostile Office file may cost.
Files are built with python-docx, openpyxl and python-pptx, or edited as zips, so
the real parsers run."""

from __future__ import annotations

import os
import zipfile
from dataclasses import replace
from io import BytesIO
from typing import Protocol

import pytest
from docx import Document
from docx import table as docx_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from api.exceptions import DocumentUnreadable
from api.services.documents.contracts import ReadContext
from api.services.documents.extraction import TRUNCATION_NOTE, extract_document
from api.services.documents.office import read_docx, read_pptx, read_xlsx
from tests.test_documents import LIMITS, make_docx, upload

SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FIRST_SHEET = "xl/worksheets/sheet1.xml"
WORD_BODY = "word/document.xml"


class Saveable(Protocol):
    """A python-docx document, an openpyxl workbook or a python-pptx deck."""

    def save(self, target: BytesIO) -> object: ...


def saved(document: Saveable) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def with_part(package: bytes, name: str, content: bytes | str) -> bytes:
    """``package`` with the part ``name`` replaced, or added when it has none."""
    source = zipfile.ZipFile(BytesIO(package))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename != name:
                target.writestr(item, source.read(item.filename))
        target.writestr(name, content)
    return buffer.getvalue()


def docx_with_body(body: str) -> bytes:
    """make_docx(), with its body replaced by ``body`` (WordprocessingML)."""
    document = f'<w:document xmlns:w="{WORD_NS}"><w:body>{body}</w:body></w:document>'
    return with_part(make_docx(), WORD_BODY, document)


def make_xlsx() -> bytes:
    workbook = Workbook()
    sales = workbook.active
    assert sales is not None
    sales.title = "Sales"
    sales.append(["Region", "Revenue", None, "Note"])
    sales.append(["North", 42, None, "two\nlines\tand a tab"])
    sales.append([])
    sales.append(["South", 3.5])
    workbook.create_sheet("Costs").append(["Rent", 1000])
    return saved(workbook)


def with_first_sheet(body: str) -> bytes:
    """make_xlsx(), with the first sheet's XML replaced by ``body`` inside <worksheet>."""
    sheet = f'<worksheet xmlns="{SPREADSHEET_NS}">{body}</worksheet>'
    return with_part(make_xlsx(), FIRST_SHEET, sheet)


def make_pptx() -> bytes:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])  # title and content
    assert first.shapes.title is not None
    first.shapes.title.text = "Quarterly results"
    first.placeholders[1].text_frame.text = "Revenue grew\nCosts fell"
    first.notes_slide.notes_text_frame.text = "Mention the new region."
    second = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    assert second.shapes.title is not None
    second.shapes.title.text = "By region"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (row, column), text in {(0, 0): "Region", (0, 1): "Revenue", (1, 0): "North"}.items():
        table.cell(row, column).text = text
    table.cell(1, 1).text = "42"
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1)).text = "Grouped words"
    deck.slides.add_slide(deck.slide_layouts[6])  # blank: no title, no text
    return saved(deck)


# --- Excel -------------------------------------------------------------------------


def test_a_workbook_gives_each_sheet_then_its_rows_as_tab_separated_values() -> None:
    result = extract_document("figures.xlsx", make_xlsx(), LIMITS)
    assert result.text.splitlines() == [
        "## Sheet: Sales",
        "Region\tRevenue\t\tNote",
        "North\t42\t\ttwo lines and a tab",
        "",
        "South\t3.5",
        "",
        "## Sheet: Costs",
        "Rent\t1000",
    ]
    assert (result.format, result.pages, result.truncated) == ("xlsx", None, False)


def test_rows_past_the_limit_are_cut_across_the_whole_workbook() -> None:
    result = extract_document("figures.xlsx", make_xlsx(), replace(LIMITS, max_spreadsheet_rows=5))
    assert "Rent" in result.text  # the fifth row read is the first of the second sheet
    assert result.truncated is False

    result = extract_document("figures.xlsx", make_xlsx(), replace(LIMITS, max_spreadsheet_rows=3))
    assert "South" not in result.text
    assert "Costs" not in result.text
    assert result.text.endswith(TRUNCATION_NOTE.strip())


def test_a_sheet_is_read_as_it_is_not_as_its_declared_size() -> None:
    # Some generators declare a wrong size; read-only openpyxl would cut to it.
    cells = '<c r="A1" t="inlineStr"><is><t>North</t></is></c><c r="B1"><v>42</v></c>'
    sheet = f'<dimension ref="A1:A1"/><sheetData><row r="1">{cells}</row></sheetData>'
    result = extract_document("figures.xlsx", with_first_sheet(sheet), LIMITS)
    assert result.text == "## Sheet: Sales\nNorth\t42\n\n## Sheet: Costs\nRent\t1000"


def test_reading_stops_once_more_cells_were_read_than_the_text_holds() -> None:
    # Each row is 16,384 cells wide (one in column XFD), however small the file.
    rows = "".join(
        f'<row r="{n}"><c r="A{n}" t="inlineStr"><is><t>row {n}</t></is></c><c r="XFD{n}"/></row>'
        for n in range(1, 101)
    )
    result = extract_document(
        "wide.xlsx", with_first_sheet(f"<sheetData>{rows}</sheetData>"), LIMITS
    )
    assert "row 1" in result.text
    assert "row 2" not in result.text
    assert result.text.endswith(TRUNCATION_NOTE.strip())


def test_a_workbook_that_unpacks_too_large_is_refused() -> None:
    with pytest.raises(DocumentUnreadable, match="Excel file is too large once unpacked"):
        extract_document("bomb.xlsx", make_xlsx(), replace(LIMITS, max_unpacked_bytes=100))


def test_something_that_is_not_a_workbook_is_unreadable() -> None:
    with pytest.raises(DocumentUnreadable, match="Excel"):
        extract_document("fake.xlsx", b"just some words", LIMITS)


# --- PowerPoint --------------------------------------------------------------------


def test_a_deck_gives_each_slide_its_title_text_tables_groups_and_notes() -> None:
    result = extract_document("deck.pptx", make_pptx(), LIMITS)
    assert result.text.splitlines() == [
        "## Slide 1: Quarterly results",
        "Revenue grew",
        "Costs fell",
        "Speaker notes: Mention the new region.",
        "",
        "## Slide 2: By region",
        "Region | Revenue",
        "North | 42",
        "Grouped words",
        "",
        "## Slide 3",
    ]
    assert (result.format, result.pages, result.truncated) == ("pptx", 3, False)


def test_slides_past_the_page_limit_are_cut_and_say_so() -> None:
    result = extract_document("deck.pptx", make_pptx(), replace(LIMITS, max_pages=1))
    assert "Slide 2" not in result.text
    assert result.pages == 3
    assert result.text.endswith(TRUNCATION_NOTE.strip())


def test_a_deck_that_unpacks_too_large_is_refused() -> None:
    with pytest.raises(DocumentUnreadable, match="PowerPoint file is too large once unpacked"):
        extract_document("bomb.pptx", make_pptx(), replace(LIMITS, max_unpacked_bytes=100))


def test_something_that_is_not_a_deck_is_unreadable() -> None:
    with pytest.raises(DocumentUnreadable, match="PowerPoint"):
        extract_document("fake.pptx", b"just some words", LIMITS)


# --- what a hostile package may cost ---------------------------------------------------

#: python-docx's own template unpacks to 0.8 MB; leave room for what a test adds.
#: Each XML part is still capped at LIMITS.max_xml_part_bytes (0.5 MB).
UNPACKED_ROOM = replace(LIMITS, max_unpacked_bytes=5_000_000)


def test_reading_a_word_file_stops_once_even_empty_paragraphs_fill_the_text() -> None:
    # 50,000 empty paragraphs: 300 kB of XML, under the part limit, then a line.
    data = docx_with_body("<w:p/>" * 50_000 + "<w:p><w:r><w:t>The end</w:t></w:r></w:p>")

    result = read_docx(data, ReadContext(UNPACKED_ROOM))

    assert result.cut_short is True
    assert len(result.text) <= LIMITS.max_chars  # not 50,000 line breaks
    with pytest.raises(DocumentUnreadable, match="No text could be read"):
        extract_document("empty.docx", data, UNPACKED_ROOM)


def test_a_word_file_with_an_xml_part_over_the_limit_is_refused_before_parsing() -> None:
    # 100,000 empty paragraphs: 600 kB of XML that add 1 kB to the file.
    data = docx_with_body("<w:p/>" * 100_000)
    assert len(data) < len(make_docx()) + 1_000

    with pytest.raises(DocumentUnreadable, match="Word file has a part too large to read"):
        extract_document("flood.docx", data, UNPACKED_ROOM)


def test_a_part_is_judged_by_its_content_not_its_name() -> None:
    disguised = b"<x>" + b"<y/>" * 150_000 + b"</x>"
    data = with_part(make_docx(), "word/media/image1.png", disguised)

    with pytest.raises(DocumentUnreadable, match="part too large"):
        extract_document("disguised.docx", data, UNPACKED_ROOM)


def test_a_picture_over_the_xml_part_limit_still_reads() -> None:
    noise = BytesIO()
    Image.frombytes("L", (800, 800), os.urandom(640_000)).save(noise, "PNG")
    assert len(noise.getvalue()) > LIMITS.max_xml_part_bytes
    document = Document()
    document.add_paragraph("Photo below")
    document.add_picture(BytesIO(noise.getvalue()))

    result = extract_document("photo.docx", saved(document), UNPACKED_ROOM)

    assert result.text == "Photo below"


def test_a_cell_spanning_millions_of_columns_is_read_once(monkeypatch: pytest.MonkeyPatch) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Wide"
    table.cell(0, 1).text = "Next"
    table.cell(0, 0)._tc.get_or_add_tcPr().grid_span = 100_000_000
    data = saved(document)

    def expanded(_row: object) -> None:
        pytest.fail("row.cells repeats a cell once per column it spans")

    # python-docx's row.cells would hand back this one cell 100 million times.
    monkeypatch.setattr(docx_table._Row, "cells", property(expanded))
    assert extract_document("span.docx", data, LIMITS).text == "Wide | Next"


def test_a_cell_merged_down_shows_its_text_on_every_row() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "North"
    for row, value in enumerate(("1", "2", "3")):
        table.cell(row, 1).text = value
    table.cell(0, 0).merge(table.cell(2, 0))

    result = extract_document("merged.docx", saved(document), LIMITS)

    assert result.text.splitlines() == ["North | 1", "North | 2", "North | 3"]


def test_reading_a_deck_stops_inside_a_long_table_once_the_text_is_full() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_table(3_000, 1, Inches(1), Inches(1), Inches(4), Inches(4))

    result = read_pptx(saved(deck), ReadContext(LIMITS))

    assert result.cut_short is True
    assert len(result.text) <= LIMITS.max_chars  # not 3,000 empty rows


def test_reading_a_workbook_stops_once_sheet_headings_fill_the_text() -> None:
    workbook = Workbook()
    for number in range(300):
        workbook.create_sheet(f"Empty {number}")

    result = read_xlsx(saved(workbook), ReadContext(LIMITS))

    assert result.cut_short is True
    assert len(result.text) < 2 * LIMITS.max_chars  # not 300 headings
    assert "Empty 299" not in result.text


# --- HTTP ----------------------------------------------------------------------------


@pytest.mark.parametrize(("name", "content"), [("a.xlsx", make_xlsx()), ("a.pptx", make_pptx())])
def test_upload_reads_office_files(auth_client: TestClient, name: str, content: bytes) -> None:
    response = upload(auth_client, name, content)
    assert response.status_code == 200
    assert response.json()["text"].startswith("## S")
