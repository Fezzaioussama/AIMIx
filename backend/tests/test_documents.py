"""Turning an uploaded document into run-input text: every format, every limit,
and the failure paths. Files are built for real, not mocked, so the parsers run."""

from __future__ import annotations

import codecs
import re
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from httpx import Response
from pypdf import PdfReader, PdfWriter
from sqlalchemy.orm import Session
from starlette.applications import Starlette
from starlette.middleware import Middleware
from starlette.requests import Request
from starlette.responses import PlainTextResponse
from starlette.routing import Route

from api.deps import app_settings, document_limits
from api.exceptions import (
    DocumentReadersBusy,
    DocumentTooLarge,
    DocumentUnreadable,
    UnsupportedDocument,
)
from api.main import create_app
from api.middleware import RequestSizeLimitMiddleware
from api.models import User
from api.services.documents.contracts import ExtractionLimits
from api.services.documents.extraction import (
    TRUNCATION_NOTE,
    extract_document,
    supported_extensions,
)
from api.services.documents.gate import extraction_slot
from api.services.documents.readers import READERS, TEXT_SUFFIXES
from config.settings import MULTIPART_OVERHEAD_BYTES, Settings, pin_settings
from tests.conftest import TEST_SETTINGS, FakeProvider
from tests.test_endpoints import make_pipeline

REPOSITORY = Path(__file__).resolve().parents[2]
LIMITS = ExtractionLimits(
    max_bytes=1_000_000,
    max_chars=1_000,
    max_pages=3,
    max_unpacked_bytes=1_000_000,
    max_xml_part_bytes=500_000,
    max_image_pixels=1_000_000,
    max_spreadsheet_rows=100,
    max_ocr_pages=3,
    ocr_max_seconds=60.0,
    ocr_dpi=72,
    max_concurrent=2,
    wait_seconds=5.0,
)


def make_pdf(pages: list[str]) -> bytes:
    """A minimal but valid PDF with one line of text per page (no parentheses in text)."""
    page_ids = [4 + 2 * index for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id, text in zip(page_ids, pages, strict=True):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page_id + 1} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode()
        objects[page_id + 1] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    body = b"%PDF-1.4\n"
    offsets: dict[int, int] = {}
    for number, content in objects.items():
        offsets[number] = len(body)
        body += f"{number} 0 obj\n".encode() + content + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{offsets[number]:010d} 00000 n \n".encode() for number in objects)
    body += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return body


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "42"
    document.add_paragraph("End of report")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- plain-text formats -------------------------------------------------------


@pytest.mark.parametrize("suffix", sorted(TEXT_SUFFIXES))
def test_every_text_format_is_read_as_text(suffix: str) -> None:
    result = extract_document(f"notes{suffix}", b"alpha\nbeta", LIMITS)
    assert result.text == "alpha\nbeta"
    assert (result.format, result.pages, result.truncated) == (suffix.lstrip("."), None, False)
    assert result.characters == len("alpha\nbeta")


def test_text_is_tidied() -> None:
    result = extract_document("a.txt", b"one  \r\n\r\n\r\n\r\ntwo\r\n", LIMITS)
    assert result.text == "one\n\ntwo"


def test_text_decoding_handles_bom_windows_1252_and_utf16() -> None:
    assert extract_document("a.txt", "\ufeffhello".encode(), LIMITS).text == "hello"
    assert extract_document("a.txt", "café".encode("cp1252"), LIMITS).text == "café"
    assert extract_document("a.txt", "héllo".encode("utf-16"), LIMITS).text == "héllo"


@pytest.mark.parametrize(
    "damage",
    [
        b"!",  # an odd byte at the end
        b"\x00\xd8",  # half a surrogate pair at the end
        b"\x00\xd8A\x00",  # half a surrogate pair inside the text
    ],
)
def test_a_damaged_utf16_file_keeps_its_readable_text(damage: bytes) -> None:
    data = codecs.BOM_UTF16_LE + "héllo".encode("utf-16-le") + damage
    assert extract_document("a.txt", data, LIMITS).text.startswith("héllo\N{REPLACEMENT CHARACTER}")


def test_a_binary_file_with_a_text_extension_is_refused() -> None:
    with pytest.raises(DocumentUnreadable, match="binary"):
        extract_document("a.txt", b"abc\x00\x01\x02def", LIMITS)


def test_html_keeps_visible_text_only() -> None:
    page = (
        b"<html><head><style>p{}</style><script>evil()</script></head>"
        b"<body><h1>Title</h1><p>Body text</p></body></html>"
    )
    assert extract_document("page.html", page, LIMITS).text == "Title\nBody text"


# --- PDF ---------------------------------------------------------------------


def test_a_pdf_gives_its_text_and_page_count() -> None:
    result = extract_document("report.pdf", make_pdf(["First page", "Second page"]), LIMITS)
    assert "First page" in result.text
    assert "Second page" in result.text
    assert (result.format, result.pages, result.truncated) == ("pdf", 2, False)


def test_a_pdf_past_the_page_limit_is_cut_and_says_so() -> None:
    pdf = make_pdf([f"Page number {n}" for n in range(1, 6)])
    result = extract_document("long.pdf", pdf, LIMITS)
    assert result.pages == 5
    assert result.truncated is True
    assert "Page number 3" in result.text
    assert "Page number 4" not in result.text
    assert result.text.endswith(TRUNCATION_NOTE.strip())


def test_text_past_the_character_limit_is_cut_with_a_note() -> None:
    result = extract_document("big.txt", b"word " * 1_000, LIMITS)
    assert result.truncated is True
    assert result.text.endswith(TRUNCATION_NOTE.strip())
    assert result.characters == len(result.text) <= LIMITS.max_chars + len(TRUNCATION_NOTE)


def test_a_password_protected_pdf_is_refused() -> None:
    writer = PdfWriter()
    for page in PdfReader(BytesIO(make_pdf(["secret"]))).pages:
        writer.add_page(page)
    writer.encrypt("hunter2")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentUnreadable, match="password"):
        extract_document("locked.pdf", buffer.getvalue(), LIMITS)


def test_a_pdf_with_no_text_layer_explains_ocr() -> None:
    with pytest.raises(DocumentUnreadable, match="OCR is not available"):
        extract_document("scan.pdf", make_pdf([""]), LIMITS)


def test_a_damaged_pdf_is_unreadable_not_a_server_error() -> None:
    with pytest.raises(DocumentUnreadable, match="PDF"):
        extract_document("broken.pdf", b"%PDF-1.4 this is not really a pdf", LIMITS)


# --- Word --------------------------------------------------------------------


def test_a_word_file_keeps_tables_between_the_paragraphs_around_them() -> None:
    result = extract_document("report.docx", make_docx(), LIMITS)
    assert result.text.splitlines() == [
        "Quarterly report",
        "Region | Revenue",
        "North | 42",
        "End of report",
    ]


def test_a_word_file_that_unpacks_too_large_is_refused() -> None:
    tight = replace(LIMITS, max_unpacked_bytes=100)
    with pytest.raises(DocumentUnreadable, match="unpacked"):
        extract_document("bomb.docx", make_docx(), tight)


def test_something_that_is_not_a_word_file_is_unreadable() -> None:
    with pytest.raises(DocumentUnreadable, match="Word"):
        extract_document("fake.docx", b"just some words", LIMITS)


# --- refusals ------------------------------------------------------------------


@pytest.mark.parametrize("name", ["tool.exe", "photo.heic", "README", "archive.tar.gz"])
def test_an_unsupported_type_names_what_is_supported(name: str) -> None:
    with pytest.raises(UnsupportedDocument, match=r"\.pdf"):
        extract_document(name, b"data", LIMITS)


def test_an_empty_file_is_refused() -> None:
    with pytest.raises(DocumentUnreadable, match="empty"):
        extract_document("a.txt", b"", LIMITS)


def test_a_file_over_the_size_limit_is_refused() -> None:
    small = replace(LIMITS, max_bytes=10)
    with pytest.raises(DocumentTooLarge, match="limit"):
        extract_document("a.txt", b"x" * 11, small)


def test_only_the_file_name_is_kept_and_the_extension_is_case_insensitive() -> None:
    result = extract_document("C:\\users\\me\\Report.TXT", b"hi", LIMITS)
    assert (result.filename, result.format) == ("Report.TXT", "txt")


def test_the_registry_lists_what_the_readers_handle() -> None:
    assert supported_extensions() == sorted(READERS)
    expected = {".pdf", ".docx", ".xlsx", ".pptx", ".html", ".txt", ".csv", ".png", ".tiff"}
    assert expected <= set(supported_extensions())


def test_the_file_picker_offers_exactly_what_the_readers_handle() -> None:
    # The frontend's DOCUMENT_FORMATS mirrors READERS; this keeps the two in step.
    source = (REPOSITORY / "frontend" / "src" / "core" / "api" / "documents.ts").read_text()
    assert set(re.findall(r"'(\.[a-z0-9]+)'", source)) == set(READERS)


# --- HTTP ----------------------------------------------------------------------


def upload(client: TestClient, name: str, content: bytes) -> Response:
    return client.post("/api/documents/extract", files={"file": (name, content)})


def test_upload_returns_the_text_in_the_documented_shape(auth_client: TestClient) -> None:
    response = upload(auth_client, "notes.txt", b"hello world")
    assert response.status_code == 200
    assert response.json() == {
        "filename": "notes.txt",
        "format": "txt",
        "text": "hello world",
        "characters": 11,
        "pages": None,
        "truncated": False,
    }


def test_upload_reads_a_real_pdf(auth_client: TestClient) -> None:
    body = upload(auth_client, "report.pdf", make_pdf(["Hello from a PDF"])).json()
    assert "Hello from a PDF" in body["text"]
    assert body["pages"] == 1


def test_upload_needs_a_signed_in_user(client: TestClient) -> None:
    response = upload(client, "notes.txt", b"hello")
    assert response.status_code == 401


def test_upload_without_a_file_is_a_validation_error(auth_client: TestClient) -> None:
    response = auth_client.post("/api/documents/extract")
    assert response.status_code == 400
    assert response.json()["error"] == "validation_error"


@pytest.mark.parametrize(
    ("name", "content", "status", "code"),
    [
        ("tool.exe", b"MZ", 415, "unsupported_document"),
        ("empty.txt", b"", 422, "document_unreadable"),
        ("scan.pdf", make_pdf([""]), 422, "document_unreadable"),
    ],
)
def test_upload_failures_use_the_error_contract(
    auth_client: TestClient, name: str, content: bytes, status: int, code: str
) -> None:
    response = upload(auth_client, name, content)
    assert response.status_code == status
    body = response.json()
    assert set(body) == {"error", "detail"}
    assert body["error"] == code


def test_upload_over_the_size_limit_is_413(app: FastAPI, auth_client: TestClient) -> None:
    app.dependency_overrides[document_limits] = lambda: replace(LIMITS, max_bytes=10)
    response = upload(auth_client, "a.txt", b"x" * 11)
    assert response.status_code == 413
    assert response.json()["error"] == "document_too_large"


def test_uploads_are_rate_limited_per_account(app: FastAPI, auth_client: TestClient) -> None:
    limited = TEST_SETTINGS.model_copy(update={"user_rate_limit_requests": 2})
    app.dependency_overrides[app_settings] = lambda: limited

    assert upload(auth_client, "a.txt", b"one").status_code == 200
    assert upload(auth_client, "b.txt", b"two").status_code == 200
    blocked = upload(auth_client, "c.txt", b"three")

    assert blocked.status_code == 429
    assert blocked.json()["error"] == "throttled"


# --- how many documents are read at once ----------------------------------------


def test_a_document_waits_for_a_free_reader_then_is_refused() -> None:
    limits = replace(LIMITS, max_concurrent=1, wait_seconds=0.05)

    busy = pytest.raises(DocumentReadersBusy, match="busy reading other documents")
    with extraction_slot(1, 60.0), busy as refusal:  # another document is being read
        extract_document("a.txt", b"hi", limits)
    assert refusal.value.headers == {"Retry-After": "1"}
    # The slot is free again once that document is done.
    assert extract_document("a.txt", b"hi", limits).text == "hi"


def test_a_busy_server_answers_503_in_the_error_contract(
    app: FastAPI, auth_client: TestClient
) -> None:
    app.dependency_overrides[document_limits] = lambda: replace(
        LIMITS, max_concurrent=1, wait_seconds=0.05
    )

    with extraction_slot(1, 60.0):
        response = upload(auth_client, "notes.txt", b"hello")

    assert response.status_code == 503
    assert response.json() == {
        "error": "documents_busy",
        "detail": "The server is busy reading other documents. Try again in a moment.",
    }
    assert response.headers["retry-after"] == "1"


def test_a_document_becomes_the_input_an_agent_reads(
    auth_client: TestClient, session: Session, user: User, fake_provider: FakeProvider
) -> None:
    pipeline = make_pipeline(session, user, "Reader")
    pdf = make_pdf(["Revenue grew by ten percent"])
    text = upload(auth_client, "report.pdf", pdf).json()["text"]

    run = auth_client.post(f"/api/pipelines/{pipeline.id}/run", json={"input": text})

    assert run.status_code == 200
    assert "Revenue grew by ten percent" in fake_provider.calls[0][0]


# --- request size guard --------------------------------------------------------


def test_the_size_guard_refuses_a_declared_oversize_body_and_passes_the_rest() -> None:
    async def echo(request: Request) -> PlainTextResponse:
        return PlainTextResponse(str(len(await request.body())))

    guarded = Starlette(
        routes=[Route("/", echo, methods=["POST"])],
        middleware=[Middleware(RequestSizeLimitMiddleware, max_bytes=10)],
    )
    client = TestClient(guarded)
    assert client.post("/", content=b"x" * 10).text == "10"
    refused = client.post("/", content=b"x" * 11)
    assert refused.status_code == 413
    assert refused.json()["error"] == "request_too_large"


def test_the_real_app_refuses_an_oversize_upload_before_reading_it() -> None:
    pin_settings(TEST_SETTINGS.model_copy(update={"document_max_bytes": 100}))
    try:
        application = create_app()
    finally:
        pin_settings(TEST_SETTINGS)
    with TestClient(application) as client:
        response = client.post(
            "/api/documents/extract", content=b"x" * (100 + MULTIPART_OVERHEAD_BYTES + 1)
        )
    assert response.status_code == 413
    assert response.json()["error"] == "request_too_large"
    assert response.headers["x-content-type-options"] == "nosniff"


def test_settings_default_to_generous_but_finite_limits() -> None:
    settings = Settings(_env_file=None)
    assert settings.document_max_bytes == 10_000_000
    assert settings.document_max_chars == 120_000
    assert settings.max_request_bytes == settings.document_max_bytes + MULTIPART_OVERHEAD_BYTES
    assert settings.document_max_xml_part_bytes == 10_000_000
    assert settings.document_max_concurrent_extractions == 2
    assert settings.document_extraction_wait_seconds == 30.0
    assert settings.document_ocr_max_seconds == 120.0
    assert settings.workspace_max_files == 2_000
