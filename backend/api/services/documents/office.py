"""Office documents: Word (.docx), Excel (.xlsx) and PowerPoint (.pptx).

All three are zip packages of XML. Before any XML is parsed, a package is refused
when it would unpack past the unpacked-size limit (zip bombs), or when it holds a
part over the XML-part limit that may be XML: python-docx, python-pptx and
openpyxl parse a whole part at once, at a cost in time and memory many times its
size. Reading then stops once the text is longer than a document may keep, with
each block's line break counted, so a million empty paragraphs stop as early as
a page of text does.
"""

from __future__ import annotations

import zipfile
from collections.abc import Iterable, Iterator
from functools import singledispatch
from io import BytesIO
from itertools import islice

from docx import Document
from docx.document import Document as WordDocument
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide
from pptx.table import _Cell as SlideCell
from pptx.text.text import TextFrame

from api.exceptions import DocumentUnreadable
from api.services.documents.contracts import ExtractionLimits, ReadContext, ReadResult

#: Between the cells of a table row, in Word and PowerPoint alike.
CELL_SEPARATOR = " | "
_W_P = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
_W_TBL = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"
#: How the media a package embeds begins: pictures, video, embedded files. A part
#: over the XML-part limit is refused unless it starts like one of these. No XML
#: parser reads past such a start, so whatever the package names or declares the
#: part to be, it cannot be XML that is slow to parse.
_MEDIA_SIGNATURES = (
    b"\x89PNG",
    b"\xff\xd8\xff",  # JPEG
    b"GIF8",
    b"BM",  # BMP
    b"II*\x00",  # TIFF, little-endian
    b"MM\x00*",  # TIFF, big-endian
    b"RIFF",  # WebP, WAV, AVI
    b"\x01\x00\x00\x00",  # EMF
    b"\xd7\xcd\xc6\x9a",  # WMF
    b"PK\x03\x04",  # an embedded package, such as the workbook behind a chart
    b"\xd0\xcf\x11\xe0",  # an embedded OLE object
    b"%PDF",
)
#: MP4 and QuickTime video name their format after a 4-byte box size.
_ISO_MEDIA = b"ftyp"
_SNIFF_BYTES = 8


def _check_package(data: bytes, limits: ExtractionLimits, kind: str) -> None:
    """Refuse a zip bomb, or a part too large to parse, before any XML is parsed."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        parts = archive.infolist()
        if sum(part.file_size for part in parts) > limits.max_unpacked_bytes:
            raise DocumentUnreadable(f"This {kind} file is too large once unpacked.")
        for part in parts:
            if part.file_size > limits.max_xml_part_bytes and not _is_media(archive, part):
                raise DocumentUnreadable(f"This {kind} file has a part too large to read.")


def _is_media(archive: zipfile.ZipFile, part: zipfile.ZipInfo) -> bool:
    with archive.open(part) as stream:
        head = stream.read(_SNIFF_BYTES)
    return head.startswith(_MEDIA_SIGNATURES) or head[4:8] == _ISO_MEDIA


def _within(blocks: Iterable[str], max_chars: int) -> tuple[list[str], bool]:
    """The first ``blocks``, until their text passes ``max_chars``, and whether any
    were left unread. Each block costs its line break too, so empty ones add up."""
    kept: list[str] = []
    length = 0
    for block in blocks:
        if length > max_chars:
            return kept, True
        kept.append(block)
        length += len(block) + 1
    return kept, False


# --- Word ------------------------------------------------------------------------


def read_docx(data: bytes, context: ReadContext) -> ReadResult:
    _check_package(data, context.limits, "Word")
    document = Document(BytesIO(data))
    blocks, more = _within(_docx_blocks(document), context.limits.max_chars)
    return ReadResult("\n".join(blocks), cut_short=more)


def _docx_blocks(document: WordDocument) -> Iterator[str]:
    """Paragraphs and table rows in body order, so a table stays between the
    paragraphs around it."""
    for child in document.element.body.iterchildren():
        if child.tag == _W_P:
            yield Paragraph(child, document).text
        elif child.tag == _W_TBL:
            yield from _table_rows(Table(child, document))


def _table_rows(table: Table) -> Iterator[str]:
    """Each row's cells. A cell merged across columns shows once; one merged down
    from the row above shows that cell's text again.

    Worked out in one pass over the XML, not through python-docx's ``row.cells``,
    which repeats a cell once per column it spans (a span of millions fits in a
    few bytes) and looks each merged-down cell up through every row above it.
    """
    above: dict[int, str] = {}  # the text last seen in each grid column
    for row in table.rows:
        column = row.grid_cols_before
        texts: list[str] = []
        for tc in row._tr.tc_lst:
            merged = tc.vMerge == "continue"
            text = above.get(column, "") if merged else _Cell(tc, table).text.strip()
            above[column] = text
            texts.append(text)
            column += tc.grid_span
        yield CELL_SEPARATOR.join(texts)


# --- Excel -----------------------------------------------------------------------


def read_xlsx(data: bytes, context: ReadContext) -> ReadResult:
    _check_package(data, context.limits, "Excel")
    # read_only streams the rows; data_only gives a formula's saved value, not its code.
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return _read_workbook(workbook, context.limits)
    finally:
        workbook.close()


def _read_workbook(workbook: Workbook, limits: ExtractionLimits) -> ReadResult:
    """Each sheet as a heading, then its rows as tab-separated values.

    Stops after ``max_spreadsheet_rows`` rows (blank ones count too) across the whole
    workbook, or once it has read more than the text may hold: a heading costs its
    length, and a row as wide as its last cell costs its width, even when empty.
    """
    lines: list[str] = []
    rows = spent = 0
    for sheet in workbook.worksheets:
        if spent > limits.max_chars:
            return ReadResult("\n".join(lines), cut_short=True)
        # Rows are padded to the size the file declares; read what is really there.
        sheet.reset_dimensions()
        lines.append(f"## Sheet: {sheet.title}")
        spent += len(lines[-1])
        for row in sheet.iter_rows(values_only=True):
            if rows == limits.max_spreadsheet_rows or spent > limits.max_chars:
                return ReadResult("\n".join(lines), cut_short=True)
            rows += 1
            spent += len(row)
            lines.append(_row_text(row))
        lines.append("")
    return ReadResult("\n".join(lines))


def _row_text(row: tuple[object, ...]) -> str:
    """Tab-separated values; a cell's own tabs and line breaks become spaces."""
    cells = ("" if value is None else " ".join(str(value).split()) for value in row)
    return "\t".join(cells).rstrip()


# --- PowerPoint ------------------------------------------------------------------


def read_pptx(data: bytes, context: ReadContext) -> ReadResult:
    limits = context.limits
    _check_package(data, limits, "PowerPoint")
    slides = Presentation(BytesIO(data)).slides
    total = len(slides)
    lines, more = _within(_deck_lines(islice(slides, limits.max_pages)), limits.max_chars)
    return ReadResult("\n".join(lines), pages=total, cut_short=more or total > limits.max_pages)


def _deck_lines(slides: Iterable[Slide]) -> Iterator[str]:
    """Every slide's lines, with a blank line between slides."""
    for number, slide in enumerate(slides, start=1):
        if number > 1:
            yield ""
        yield from _slide_lines(number, slide)


def _slide_lines(number: int, slide: Slide) -> Iterator[str]:
    """'## Slide N: <title>', then the slide's other text, then its speaker notes."""
    title = slide.shapes.title
    heading = f"## Slide {number}"
    if title is not None and title.text.strip():
        heading += f": {' '.join(title.text.split())}"
    yield heading
    for shape in slide.shapes:
        if shape != title:
            yield from _shape_lines(shape)
    yield from _notes(slide)


def _notes(slide: Slide) -> list[str]:
    # Reading notes_slide would create one, so ask first.
    frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
    text = _text(frame) if frame is not None else ""
    return [f"Speaker notes: {text}"] if text else []


def _text(frame: TextFrame) -> str:
    # python-pptx writes a line break inside a paragraph as a vertical tab.
    return frame.text.replace("\v", "\n").strip()


# The text a shape shows, by kind of shape (§7 Strategy, through singledispatch):
# a new kind of shape is one more registered function. Each yields its lines one
# at a time, so reading stops inside a huge group or table once the text is full.
@singledispatch
def _shape_lines(shape: BaseShape) -> Iterator[str]:
    """Pictures, charts, connectors and other shapes without text show none."""
    return iter(())


@_shape_lines.register
def _group_lines(shape: GroupShape) -> Iterator[str]:
    for child in shape.shapes:
        yield from _shape_lines(child)


@_shape_lines.register
def _table_lines(shape: GraphicFrame) -> Iterator[str]:
    if not shape.has_table:
        return
    # From the row list itself: indexing table.rows looks the whole list up again
    # for every row, which takes time in the square of their number.
    for tr in shape.table._tbl.tr_lst:
        yield CELL_SEPARATOR.join(_text(SlideCell(tc, shape).text_frame) for tc in tr.tc_lst)


@_shape_lines.register
def _text_lines(shape: Shape) -> Iterator[str]:
    if shape.has_text_frame:
        yield _text(shape.text_frame)
