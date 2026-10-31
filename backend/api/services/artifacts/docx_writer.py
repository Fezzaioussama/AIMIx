"""Markdown blocks into a Word document with python-docx.

One writer per block type, chosen by ``functools.singledispatch`` (§7 Strategy +
Registry). Word's own styles carry the look (Heading 1-6, List Bullet, List
Number, Quote, Table Grid), so the file stays as editable as any other.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from functools import singledispatch
from io import BytesIO

from docx import Document
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as WordParagraph

from api.services.artifacts.contracts import AI_GENERATED_KEYWORD, AI_GENERATED_NOTE, RenderOptions
from api.services.artifacts.markdown import (
    Block,
    CodeBlock,
    Heading,
    ListBlock,
    Paragraph,
    Quote,
    Rule,
    Spans,
    Table,
    parse_markdown,
    web_link,
)

#: The template's list styles go three levels deep ("List Bullet 3").
LIST_LEVELS = 3
LIST_INDENT = Pt(18)
CODE_FONT = "Courier New"
CODE_SIZE = Pt(9.5)
LINK_COLOR = RGBColor(0x1D, 0x4E, 0xD8)
RULE_COLOR = "A0AEC0"


@dataclass(frozen=True)
class _Place:
    """Where a block lands: its list depth, and the style and numbering its
    paragraph takes (a list item's first paragraph, a quote, ...)."""

    depth: int = 0
    style: str | None = None
    numbering: int | None = None


def render_docx(content: str, options: RenderOptions) -> bytes:
    """The Markdown ``content`` as a .docx file. ``options`` concern PDFs only."""
    document = Document()
    properties = document.core_properties
    properties.author = "AIMIx"
    properties.comments = AI_GENERATED_NOTE
    properties.keywords = AI_GENERATED_KEYWORD
    for block in parse_markdown(content):
        _write(block, document, _Place())
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@singledispatch
def _write(block: Block, document: WordDocument, place: _Place) -> None:
    raise TypeError(f"No Word writer for {type(block).__name__}.")


@_write.register
def _heading(block: Heading, document: WordDocument, place: _Place) -> None:
    _add_spans(document.add_heading(level=block.level), block.spans)


@_write.register
def _paragraph(block: Paragraph, document: WordDocument, place: _Place) -> None:
    paragraph = document.add_paragraph(style=place.style)
    if place.numbering is not None:
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_numId().val = place.numbering
        numbering.get_or_add_ilvl().val = 0
    _add_spans(paragraph, block.spans)


@_write.register
def _list(block: ListBlock, document: WordDocument, place: _Place) -> None:
    level = min(place.depth + 1, LIST_LEVELS)
    suffix = "" if level == 1 else f" {level}"
    style = f"List {'Number' if block.ordered else 'Bullet'}{suffix}"
    numbering = _restart_numbering(document, style, block.start) if block.ordered else None
    item_place = _Place(level, style, numbering)
    # Paragraphs after an item's first continue it, indented, without a bullet.
    rest_place = _Place(level, f"List Continue{suffix}")
    for item in block.items:
        for index, child in enumerate(item):
            _write(child, document, rest_place if index else item_place)


@_write.register
def _quote(block: Quote, document: WordDocument, place: _Place) -> None:
    for child in block.blocks:
        _write(child, document, replace(place, style="Quote", numbering=None))


@_write.register
def _code(block: CodeBlock, document: WordDocument, place: _Place) -> None:
    # Never a list style: that would number or bullet the code.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = LIST_INDENT * place.depth or None
    run = paragraph.add_run(block.text)
    run.font.name = CODE_FONT
    run.font.size = CODE_SIZE


@_write.register
def _table(block: Table, document: WordDocument, place: _Place) -> None:
    rows = [block.header, *block.rows] if block.header else list(block.rows)
    columns = max((len(row) for row in rows), default=0)
    if not columns:
        return
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for index, (word_row, row) in enumerate(zip(table.rows, rows, strict=True)):
        # A short row leaves its last cells empty.
        for cell, spans in zip(word_row.cells, row, strict=False):
            _add_spans(cell.paragraphs[0], spans, bold=bool(block.header) and index == 0)


@_write.register
def _rule(block: Rule, document: WordDocument, place: _Place) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for name, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", RULE_COLOR)):
        bottom.set(qn(f"w:{name}"), value)
    border.append(bottom)
    document.add_paragraph()._p.get_or_add_pPr().append(border)


def _restart_numbering(document: WordDocument, style: str, start: int) -> int:
    """A numbering of its own for one ordered list, counting from ``start``,
    so it does not carry on from the list before it."""
    numbering = document.part.numbering_part.element
    style_numbering = document.styles[style].element.pPr.numPr.numId.val
    instance = numbering.add_num(numbering.num_having_numId(style_numbering).abstractNumId.val)
    instance.add_lvlOverride(ilvl=0).add_startOverride(start)
    return int(instance.numId)


def _add_spans(paragraph: WordParagraph, spans: Spans, *, bold: bool = False) -> None:
    for span in spans:
        run = paragraph.add_run(span.text)
        run.bold = bold or span.bold or None
        run.italic = span.italic or None
        if span.code:
            run.font.name = CODE_FONT
        target = web_link(span)
        if target:
            _link(paragraph, run._r, target)
            run.font.underline = True
            run.font.color.rgb = LINK_COLOR


def _link(paragraph: WordParagraph, run: object, target: str) -> None:
    """Wrap a run that was just added in a hyperlink to ``target``."""
    relation = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    paragraph._p.append(hyperlink)
    hyperlink.append(run)
